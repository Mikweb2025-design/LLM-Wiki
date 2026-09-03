"""Processamento documenti per estrazione testo"""
import os
import threading
import time
from pathlib import Path

# Cache del testo estratto. Key = path, value = (mtime, ts, text)
# Invalidazione: se mtime del file cambia, ricomputiamo.
# TTL hard: 30 min — protegge da file rimossi senza notifica.
_PROC_CACHE: dict = {}
_PROC_CACHE_LOCK = threading.Lock()
_PROC_CACHE_TTL = 1800.0
_PROC_CACHE_MAX = 256


def _cache_get(file_path: str):
    try:
        mtime = os.path.getmtime(file_path)
    except OSError:
        return None
    with _PROC_CACHE_LOCK:
        hit = _PROC_CACHE.get(file_path)
        if hit and hit[0] == mtime and (time.monotonic() - hit[1]) < _PROC_CACHE_TTL:
            return hit[2]
    return None


def _cache_put(file_path: str, text: str):
    try:
        mtime = os.path.getmtime(file_path)
    except OSError:
        return
    with _PROC_CACHE_LOCK:
        if len(_PROC_CACHE) >= _PROC_CACHE_MAX:
            # rimuovi l'entry più vecchia
            oldest = min(_PROC_CACHE, key=lambda k: _PROC_CACHE[k][1])
            _PROC_CACHE.pop(oldest, None)
        _PROC_CACHE[file_path] = (mtime, time.monotonic(), text)


def invalidate_processor_cache() -> None:
    with _PROC_CACHE_LOCK:
        _PROC_CACHE.clear()


def _ionos_vision_ocr(image_path: str) -> str:
    """Prova OCR via IONOS Vision (image → text). Ritorna stringa vuota se fallisce."""
    try:
        from app.config import IONOS_API_KEY, IONOS_BASE_URL, IONOS_VISION_MODEL, OCR_HYBRID_ENABLED
        if not OCR_HYBRID_ENABLED or not IONOS_API_KEY:
            return ""
        import base64, requests
        with open(image_path, "rb") as f:
            b64 = base64.b64encode(f.read()).decode()
        # Determina mime
        ext = Path(image_path).suffix.lower()
        mime = "image/png" if ext==".png" else "image/jpeg"
        url = f"{IONOS_BASE_URL}/chat/completions"
        headers = {"Authorization": f"Bearer {IONOS_API_KEY}", "Content-Type": "application/json"}
        payload = {
            "model": IONOS_VISION_MODEL,
            "messages": [
                {"role": "system", "content": "Sei un OCR. Trascrivi fedelmente tutto il testo visibile nell'immagine, mantieni righe e numeri. Rispondi solo con il testo trascritto, senza commenti."},
                {"role": "user", "content": [
                    {"type": "text", "text": "Trascrivi questo documento:"},
                    {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}}
                ]}
            ],
            "temperature": 0.0,
            "max_tokens": 3000,
        }
        resp = requests.post(url, headers=headers, json=payload, timeout=60)
        if resp.status_code==200:
            j=resp.json()
            txt=j["choices"][0]["message"]["content"]
            return (txt or "").strip()
        else:
            # fallback: log ma non crash
            print(f"[WARN] IONOS Vision OCR {resp.status_code}: {resp.text[:200]}")
            return ""
    except Exception as e:
        print(f"[WARN] IONOS Vision exception: {e}")
        return ""

def extract_text_from_pdf(file_path: str) -> str:
    """Estrae testo da PDF — Hybrid: pypdf text → Tesseract per pagina → IONOS Vision fallback per pagine povere."""
    try:
        from pypdf import PdfReader
        from app.config import OCR_MIN_CHARS_PER_PAGE
        reader = PdfReader(file_path)
        page_texts = []
        has_text_pages = 0
        for idx, page in enumerate(reader.pages):
            txt = (page.extract_text() or "").strip()
            if txt and len(txt) >= 20:
                has_text_pages += 1
            page_texts.append(txt)

        # Se almeno metà pagine hanno testo, usa estrazione classica (mantiene page break)
        if has_text_pages >= len(page_texts) / 2 and has_text_pages > 0:
            text = "\n".join(t for t in page_texts if t)
            # Ma per pagine vuote prova comunque OCR locale + Vision
            if any(len(t.strip()) < 10 for t in page_texts):
                # OCR solo per pagine vuote
                try:
                    import pytesseract, tempfile
                    from pdf2image import convert_from_path
                    images = convert_from_path(file_path, dpi=200)
                    with tempfile.TemporaryDirectory(prefix="llmwiki_ocr_") as td:
                        enriched=[]
                        for i, t in enumerate(page_texts):
                            if t and len(t.strip()) >= 20:
                                enriched.append(t)
                            else:
                                img = images[i] if i < len(images) else None
                                if img is None:
                                    enriched.append(t)
                                    continue
                                png = os.path.join(td, f"p_{i}.png")
                                img.save(png, "PNG")
                                ocr = pytesseract.image_to_string(png, lang='ita+eng', config='--oem 1 --psm 6').strip()
                                if len(ocr) < OCR_MIN_CHARS_PER_PAGE:
                                    vision = _ionos_vision_ocr(png)
                                    if len(vision) > len(ocr):
                                        ocr = vision
                                enriched.append(ocr if ocr else t)
                        text = "\n".join(e for e in enriched if e)
                except Exception as e:
                    print(f"[WARN] OCR ibrido pagine vuote: {e}")
            return text.strip() if text.strip() else ""

        # Nessun testo utile → OCR completo ibrido pagina-per-pagina
        try:
            import pytesseract, tempfile
            from pdf2image import convert_from_path
            from app.config import OCR_MIN_CHARS_PER_PAGE
            images = convert_from_path(file_path, dpi=200)
            with tempfile.TemporaryDirectory(prefix="llmwiki_ocr_") as td:
                out_parts=[]
                for i, image in enumerate(images):
                    png_path = os.path.join(td, f"page_{i}.png")
                    image.save(png_path, "PNG")
                    # Prima Tesseract locale
                    try:
                        ocr_text = pytesseract.image_to_string(png_path, lang='ita+eng', config='--oem 1 --psm 6').strip()
                    except Exception as inner:
                        print(f"[WARN] OCR Tesseract pagina {i}: {inner}")
                        ocr_text=""
                    # Se povero, prova IONOS Vision
                    if len(ocr_text) < OCR_MIN_CHARS_PER_PAGE:
                        vision = _ionos_vision_ocr(png_path)
                        if vision and len(vision.strip()) > len(ocr_text):
                            ocr_text = vision
                            print(f"[INFO] Pagina {i+1} usata IONOS Vision ({len(vision)} chars)")
                    if ocr_text:
                        out_parts.append(ocr_text)
                text = "\n".join(out_parts)
                if text.strip():
                    return text.strip()
        except Exception as e:
            print(f"[WARN] OCR ibrido completo fallito ({file_path}): {e}")

        # Ultimo fallback: restituisci quello che pypdf aveva anche se scarso
        fallback = "\n".join(t for t in page_texts if t)
        return fallback.strip() if fallback.strip() else ""
    except Exception as e:
        raise ValueError(f"Errore lettura PDF: {str(e)}")


def get_pdf_pages(file_path: str) -> list:
    """Ritorna lista pagine [{page:int, text:str, source:str}] per Citations 2.0. Source: text|ocr_tesseract|ocr_ionos"""
    try:
        from pypdf import PdfReader
        reader = PdfReader(file_path)
        pages=[]
        for idx, page in enumerate(reader.pages, 1):
            t=(page.extract_text() or "").strip()
            if t and len(t)>=20:
                pages.append({"page": idx, "text": t, "source": "text"})
            else:
                # placeholder — fill via OCR hybrid path in extract_text_from_pdf
                # per semplicità, se nessuna text, segna come ocr_needed e lascia text vuoto
                # Il chiamante userà extract_text_from_pdf per il full text, ma pages mantiene page number
                pages.append({"page": idx, "text": t, "source": "text" if t else "ocr_needed"})
        # Se ci sono ocr_needed, prova a popolare con OCR (riusa extract logic ma per-page)
        if any(p["source"]=="ocr_needed" for p in pages):
            full = extract_text_from_pdf(file_path)
            # stima split per pagina: se full contiene ~ uguale per pagine ocr, ridistribuisci uniforme? Semplice: lascia full come unico text per ora
            # Meglio: ricalcola con pdf2image se disponibile
            try:
                import pytesseract, tempfile
                from pdf2image import convert_from_path
                from app.config import OCR_MIN_CHARS_PER_PAGE
                images = convert_from_path(file_path, dpi=180)
                with tempfile.TemporaryDirectory(prefix="llmwiki_pages_") as td:
                    for i, p in enumerate(pages):
                        if p["source"]!="ocr_needed":
                            continue
                        if i >= len(images): continue
                        png = os.path.join(td, f"pg_{i}.png")
                        images[i].save(png, "PNG")
                        ocr = pytesseract.image_to_string(png, lang='ita+eng', config='--oem 1 --psm 6').strip()
                        src="ocr_tesseract"
                        if len(ocr) < OCR_MIN_CHARS_PER_PAGE:
                            vision=_ionos_vision_ocr(png)
                            if vision and len(vision)>len(ocr):
                                ocr=vision; src="ocr_ionos"
                        pages[i]["text"]=ocr
                        pages[i]["source"]=src
            except Exception as e:
                print(f"[WARN] get_pdf_pages OCR fill: {e}")
        return pages
    except Exception as e:
        print(f"[WARN] get_pdf_pages: {e}")
        return []


def extract_text_from_excel(file_path: str) -> str:
    """Estrae dati da file Excel (ottimizzato: read_only + streaming)."""
    try:
        import openpyxl
        text_parts = []
        # read_only=True usa iteratore streaming, molto più veloce su file grandi
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            text_parts.append(f"# Foglio: {sheet_name}")
            empty_rows = 0
            for row in ws.iter_rows(values_only=True):
                row_cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                if not row_cells:
                    empty_rows += 1
                    if empty_rows > 20:  # stop su fogli con tante righe vuote finali
                        break
                    continue
                empty_rows = 0
                text_parts.append(" | ".join(row_cells))
                if len(text_parts) > 5000:  # guard per fogli enormi
                    text_parts.append("[... foglio troncato per performance ...]")
                    break
        try:
            wb.close()
        except Exception:
            pass
        return "\n".join(text_parts).strip()
    except Exception as e:
        raise ValueError(f"Impossibile leggere il file Excel: {str(e)}")


def extract_text_from_pptx(file_path: str) -> str:
    """Estrae testo da PowerPoint (slide + tabelle + note)."""
    try:
        from pptx import Presentation
        prs = Presentation(file_path)
        parts = []
        for idx, slide in enumerate(prs.slides, 1):
            parts.append(f"--- Slide {idx} ---")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    txt = shape.text.strip()
                    if txt:
                        parts.append(txt)
                if shape.has_table:
                    for row in shape.table.rows:
                        cells = [c.text.strip() for c in row.cells if c.text.strip()]
                        if cells:
                            parts.append(" | ".join(cells))
            # note del relatore
            try:
                if slide.notes_slide and slide.notes_slide.placeholders:
                    for ph in slide.notes_slide.placeholders:
                        if ph.has_text_frame and ph.text.strip():
                            parts.append(f"[Note] {ph.text.strip()}")
            except Exception:
                pass
        return "\n".join(parts).strip()
    except Exception as e:
        raise ValueError(f"Errore lettura PPTX: {str(e)}")


def extract_text_from_csv(file_path: str) -> str:
    """Estrae dati da CSV con auto-detect delimiter."""
    try:
        import csv
        # prova a sniffare il delimitatore
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            sample = f.read(4096)
            f.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=[',', ';', '\t', '|'])
            except Exception:
                dialect = csv.excel
            reader = csv.reader(f, dialect)
            rows = []
            for i, row in enumerate(reader):
                if i > 5000:
                    rows.append("[... CSV troncato a 5000 righe ...]")
                    break
                cells = [c.strip() for c in row if c.strip()]
                if cells:
                    rows.append(" | ".join(cells))
        return "\n".join(rows).strip()
    except Exception as e:
        raise ValueError(f"Errore lettura CSV: {str(e)}")


def extract_text_from_html(file_path: str) -> str:
    """Estrae testo pulito da HTML/MD."""
    try:
        from bs4 import BeautifulSoup
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            raw = f.read()
        # se è markdown puro, ritorna direttamente
        if file_path.lower().endswith('.md'):
            return raw.strip()
        soup = BeautifulSoup(raw, 'html.parser')
        # rimuovi script/style
        for tag in soup(["script", "style", "noscript"]):
            tag.decompose()
        text = soup.get_text(separator="\n")
        # normalizza righe vuote
        lines = [l.strip() for l in text.splitlines()]
        lines = [l for l in lines if l]
        return "\n".join(lines).strip()
    except Exception as e:
        # fallback: lettura raw
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read().strip()
        except Exception:
            raise ValueError(f"Errore lettura HTML: {str(e)}")


def extract_text_from_docx(file_path: str) -> str:
    """Estrae testo da file Word"""
    try:
        from docx import Document
        doc = Document(file_path)
        text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
        return text.strip()
    except Exception as e:
        raise ValueError(f"Errore lettura Word: {str(e)}")


def extract_text_from_image(file_path: str) -> str:
    """OCR su immagini singole (PNG/JPG/...). Ritorna stringa vuota se OCR non disponibile.
    Passa il path direttamente a pytesseract — non l'oggetto PIL — per evitare
    problemi con TMPDIR sandboxati."""
    try:
        import pytesseract
        # config OCR: --oem 1 (LSTM) + --psm 6 (blocco uniforme) è buon compromesso
        return pytesseract.image_to_string(file_path, lang='ita+eng', config='--oem 1 --psm 6').strip()
    except Exception as e:
        print(f"[WARN] OCR immagine fallito ({file_path}): {e}")
        return ""


def extract_text_from_rtf(file_path: str) -> str:
    """Estrae testo da RTF (fallback: regex strip)."""
    try:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            raw = f.read()
        # prova con striprtf se disponibile
        try:
            from striprtf.striprtf import rtf_to_text
            return rtf_to_text(raw).strip()
        except ImportError:
            pass
        # fallback naive: rimuovi comandi RTF
        import re
        text = re.sub(r'\\[a-z]+\d*\s?', ' ', raw)
        text = re.sub(r'[{}]', ' ', text)
        text = re.sub(r'\s+', ' ', text)
        return text.strip()[:50000]
    except Exception as e:
        raise ValueError(f"Errore lettura RTF: {str(e)}")


def process_document(file_path: str) -> str:
    """Processa diversi tipi di file. Cache TTL + invalidazione su mtime."""
    cached = _cache_get(file_path)
    if cached is not None:
        return cached

    ext = Path(file_path).suffix.lower()

    if ext in [".xlsx", ".xls"]:
        text = extract_text_from_excel(file_path)
    elif ext == ".pdf":
        text = extract_text_from_pdf(file_path)
    elif ext == ".docx":
        text = extract_text_from_docx(file_path)
    elif ext in [".pptx", ".ppt"]:
        text = extract_text_from_pptx(file_path)
    elif ext == ".csv":
        text = extract_text_from_csv(file_path)
    elif ext in [".html", ".htm", ".md"]:
        text = extract_text_from_html(file_path)
    elif ext == ".rtf":
        text = extract_text_from_rtf(file_path)
    elif ext == ".txt":
        with open(file_path, 'r', encoding='utf-8') as f:
            text = f.read()
    elif ext in [".png", ".jpg", ".jpeg", ".gif", ".tiff", ".bmp", ".webp"]:
        # OCR diretto su immagini (prima il content endpoint le leggeva come testo
        # binario sporco, da qui il "Confronta" non funzionava su WhatsApp-screenshot)
        text = extract_text_from_image(file_path)
    else:
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                text = f.read()
        except Exception:
            raise ValueError(f"Formato non supportato: {ext}")

    _cache_put(file_path, text)
    return text



def get_document_metadata(file_path: str) -> dict:
    """Ottiene metadati documento"""
    path = Path(file_path)
    
    return {
        "filename": path.name,
        "extension": path.suffix.lower(),
        "size_bytes": path.stat().st_size,
        "modified": path.stat().st_mtime,
        "created": path.stat().st_ctime,
        "dir": str(path.parent)
    }


def scan_documents_directory(directory: str = None) -> list:
    """Scansiona cartella per documenti supportati"""
    from app.config import DOCUMENTS_DIR
    
    directory = directory or str(DOCUMENTS_DIR)
    
    supported_extensions = {
        '.pdf', '.docx', '.xlsx', '.xls', '.txt',
        '.jpg', '.jpeg', '.png', '.gif', '.tiff', '.bmp',
        '.rtf', '.html', '.md', '.csv'
    }
    
    files = []
    for root, _, filenames in os.walk(directory):
        for filename in filenames:
            path = Path(root) / filename
            ext = path.suffix.lower()
            if ext in supported_extensions:
                files.append(str(path))
    
    return files


def scan_custom_directory(directory: str) -> list:
    """Scansiona cartella personalizzata"""
    if not directory:
        raise ValueError("Directory mancante")
    
    return scan_documents_directory(directory)
